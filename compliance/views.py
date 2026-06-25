import io

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect, render
from docx import Document

from .models import ControlDomain, ControlItem


# =================================================================
# ISMS Compliance Helpers
# =================================================================

IT_DEPARTMENT = "IT"


def _normalize_department(department: str) -> str:
    return (department or "").strip()


def can_view_compliance(user) -> bool:
    """
    Only IT Department users may access Compliance / Controls / Export.
    UI hiding is not enough; this is the backend guard.
    """
    if not getattr(user, "is_authenticated", False):
        return False

    profile = getattr(user, "profile", None)
    if not profile:
        return False

    return _normalize_department(profile.department).lower() == IT_DEPARTMENT.lower()


def compliance_required(view_func):
    def _wrapped(request, *args, **kwargs):
        if not can_view_compliance(request.user):
            messages.error(request, "You do not have permission to access Compliance features.")
            return redirect("portal_dashboard")
        return view_func(request, *args, **kwargs)
    return _wrapped


def get_compliance_summary():
    """
    Shared compliance summary for Compliance Dashboard.
    Keeps template variables consistent:
    - domain_stats
    - control_domain_count
    - control_total
    - compliant_total
    - compliance_rate
    """
    domains = ControlDomain.objects.prefetch_related("items__evidences").all()

    domain_stats = []
    control_total = 0
    compliant_total = 0

    for domain in domains:
        items = list(domain.items.all())
        total = len(items)
        compliant = sum(
            1 for item in items
            if (ev := item.evidences.first()) and ev.status
        )

        control_total += total
        compliant_total += compliant

        domain_stats.append({
            "domain": domain,
            "total": total,
            "compliant": compliant,
            "rate": round(compliant / total * 100, 1) if total else 0,
        })

    compliance_rate = round(compliant_total / control_total * 100, 1) if control_total else 0

    return {
        "domain_stats": domain_stats,
        "control_domain_count": len(domain_stats),
        "control_total": control_total,
        "compliant_total": compliant_total,
        "compliance_rate": compliance_rate,
    }


# =================================================================
# ISMS Compliance Views
# =================================================================

@login_required
@compliance_required
def compliance_dashboard(request):
    """
    ISMS Compliance Dashboard.
    IT Department only.
    """
    context = get_compliance_summary()
    return render(request, "compliance/dashboard.html", context)


@login_required
@compliance_required
def control_list(request):
    """
    ISO 27001:2022 Annex A control list.
    IT Department only.
    """
    domains = ControlDomain.objects.prefetch_related("items__evidences").all()
    return render(request, "compliance/control_list.html", {"domains": domains})


@login_required
@compliance_required
def export_report_docx(request):
    """
    Export ISO 27001 compliance status Word report.
    IT Department only.
    """
    doc = Document()
    doc.add_heading("USCIP - ISO 27001:2022 Compliance Report", 0)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["控制項 ID", "標題", "風險等級", "合規狀態"]

    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header

    for item in ControlItem.objects.select_related("domain").prefetch_related("evidences").all():
        latest_ev = item.evidences.first()
        row = table.add_row().cells
        row[0].text = item.item_code
        row[1].text = item.title
        row[2].text = item.get_risk_level_display()
        row[3].text = "Compliant" if latest_ev and latest_ev.status else "Pending"

    file_obj = io.BytesIO()
    doc.save(file_obj)
    file_obj.seek(0)

    response = HttpResponse(
        file_obj.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response["Content-Disposition"] = "attachment; filename=USCIP_Compliance_Report.docx"
    return response
